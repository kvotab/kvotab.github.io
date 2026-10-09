"""End-to-end checks of skb_qa_summary.html with folders of workbooks.

A tree of small workbooks is built in a temporary folder, then dropped on the
page, chosen with its folder picker and dropped again, each through Chrome's
own drag and file-chooser machinery (CDP), so the page gets trusted events
and real directory entries. Checked:

- every QC row of every workbook in the tree, with its folder, file, sheet
  and Excel row (a sheet whose header is in row 3 included);
- .xlsm, .xls and a QA-headed sheet are read; Excel's "~$" owner files, "._"
  copies, hidden folders and other files are passed over;
- two workbooks of one name in two folders are two files;
- the Folders filter lists the tree and a folder takes its subfolders along;
- the CSV export starts with the Folder column;
- a folder read again replaces its rows instead of doubling them;
- a workbook added on its own shows no folder;
- a workbook that cannot be read, a folder with no workbook, Stop and
  Clear all in the middle of a read;
- a drop outside the drop zone is taken, and so is a dropped item with no
  entry (a mail attachment, say);
- "Choose a folder", clicked or by Enter, opens the folder chooser and
  nothing else;
- the SheetJS loaded is the vendored copy, 0.20.2 or later;
- no sideways scroll at 1280 px or on a 375 px phone, whose fields are 16 px
  (smaller ones make iOS zoom), and the summary cards read in the dark theme.

Serve the repository root and start headless Chrome, by default on ports 8765
and 9222 (QC_HTTP_PORT and QC_CDP_PORT choose others):

    python3 -m http.server 8765
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome" --headless=new \\
        --remote-debugging-port=9222 --user-data-dir=<fresh dir> about:blank
    python3 resources/tests/skb_qa_summary/test-ui.py

Needs openpyxl and websockets. Exit status is 0 when every check passes.
"""
import asyncio
import base64
import csv
import io
import json
import os
import shutil
import sys
import tempfile
import urllib.request

import websockets
from openpyxl import Workbook

HTTP = int(os.environ.get('QC_HTTP_PORT', '8765'))
CDP = int(os.environ.get('QC_CDP_PORT', '9222'))
URL = f'http://127.0.0.1:{HTTP}/skb_qa_summary.html'

ROOT = 'QC granskning'
QC_HEAD = ['ID', 'Value', 'QC reviewer', 'QC comment', 'QC status']


def book(path, sheets):
    """A workbook at path; sheets maps a sheet name to its rows, top first (None: an empty row)."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r, row in enumerate(rows, 1):
            for c, value in enumerate(row or [], 1):
                if value is not None:
                    ws.cell(row=r, column=c, value=value)
    wb.save(path)


def junk(path, data=b'not a workbook at all'):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'wb') as f:
        f.write(data)


def build_tree(base, xls_bytes):
    """The fixture folders, and the rows the page must find in ROOT."""
    root = os.path.join(base, ROOT)
    data_sheet = [['a', 'b', 'c']] + [[i, i * 2, i * 3] for i in range(50)]
    book(os.path.join(root, 'PSAR 2024', 'params.xlsx'), {
        'Params': [QC_HEAD] + [['p1', 1, 'revA', '', 'OK-PSAR']] * 3 + [['p2', 2, 'revB', 'check unit', 'Preliminary']] * 2,
        'Data': data_sheet,
    })
    book(os.path.join(root, 'PSAR 2025', 'params.xlsx'), {
        'Data': data_sheet,
        'Params': [QC_HEAD, ['p1', 1, 'revA', '', 'OK-FSAR']] + [['p3', 3, 'revC', '', 'In review']] * 2,
    })
    book(os.path.join(root, 'PSAR 2025', 'äldre', 'legacy.xlsm'), {
        'QA': [['ID', 'QA reviewer', 'QA comment', 'QA status'], ['p4', 'revD', 'old heading', 'OK-UPSAR']],
    })
    # The header in row 3, rows 1 and 2 empty: the sheet's range starts there.
    book(os.path.join(root, 'PSAR 2025', 'offset.xlsx'), {
        'Lower': [None, None, ['ID', 'QC reviewer', 'QC comment', 'QC status'],
                  ['p6', 'revA', '', 'OK-PSAR'], ['p7', 'revA', '', 'Preliminary']],
    })
    junk(os.path.join(root, 'PSAR 2025', 'gammal.xls'), xls_bytes)
    book(os.path.join(root, 'notes.xlsx'), {'Notes': [['Topic', 'Text'], ['x', 'y']]})
    junk(os.path.join(root, '~$params.xlsx'), b'\x05owner' + b'\x00' * 150)
    junk(os.path.join(root, '._notes.xlsx'), b'\x00\x05\x16\x07' + b'\x00' * 60)
    book(os.path.join(root, '.hidden', 'secret.xlsx'), {'Params': [QC_HEAD, ['p99', 0, 'revZ', '', 'OK-PSAR']]})
    junk(os.path.join(root, 'readme.txt'), b'plain text')

    book(os.path.join(base, 'Trasig', 'ok.xlsx'), {'Params': [QC_HEAD, ['p8', 8, 'revA', '', 'OK-PSAR']]})
    # A workbook cut short. Plain text would not do: SheetJS reads it as a one-cell sheet.
    with open(os.path.join(base, 'Trasig', 'ok.xlsx'), 'rb') as f:
        junk(os.path.join(base, 'Trasig', 'broken.xlsx'), f.read()[:2000])
    junk(os.path.join(base, 'Tom', 'readme.txt'), b'nothing here')
    book(os.path.join(base, 'Lös', 'params.xlsx'), {'Params': [QC_HEAD, ['p1', 1, 'revE', '', 'OK-FSAR']]})
    for k in range(1, 13):
        book(os.path.join(base, 'Många', f'book{k:02}.xlsx'), {'Params': [QC_HEAD, [f'm{k:02}', k, 'revM', '', 'OK-PSAR']]})

    # In the order the page reads them: a folder at a time, "äldre" sorting as "aldre".
    f24, f25 = f'{ROOT}/PSAR 2024', f'{ROOT}/PSAR 2025'
    expected = (
        [(f24, 'params.xlsx', 'Params', r, 'p1', 'revA', 'OK-PSAR', '') for r in (2, 3, 4)]
        + [(f24, 'params.xlsx', 'Params', r, 'p2', 'revB', 'Preliminary', 'check unit') for r in (5, 6)]
        + [(f'{f25}/äldre', 'legacy.xlsm', 'QA', 2, 'p4', 'revD', 'OK-UPSAR', 'old heading')]
        + [(f25, 'gammal.xls', 'Gammal', 2, 'p5', 'revB', 'OK-PSAR', 'old format')]
        + [(f25, 'offset.xlsx', 'Lower', 4, 'p6', 'revA', 'OK-PSAR', ''), (f25, 'offset.xlsx', 'Lower', 5, 'p7', 'revA', 'Preliminary', '')]
        + [(f25, 'params.xlsx', 'Params', 2, 'p1', 'revA', 'OK-FSAR', '')]
        + [(f25, 'params.xlsx', 'Params', r, 'p3', 'revC', 'In review', '') for r in (3, 4)]
    )
    return expected


class Page:
    def __init__(self, ws):
        self.ws = ws
        self.n = 0
        self.events = []
        self.errors = []

    async def send(self, method, params=None):
        self.n += 1
        await self.ws.send(json.dumps({'id': self.n, 'method': method, 'params': params or {}}))
        while True:
            r = json.loads(await self.ws.recv())
            if r.get('id') == self.n:
                if 'error' in r:
                    raise RuntimeError(f'{method}: {r["error"]}')
                return r.get('result', {})
            self.note(r)

    def note(self, r):
        m = r.get('method')
        if m == 'Runtime.exceptionThrown':
            d = r['params']['exceptionDetails']
            self.errors.append((d.get('exception', {}).get('description') or d.get('text', ''))[:300])
        elif m == 'Runtime.consoleAPICalled' and r['params']['type'] == 'error':
            self.errors.append(' '.join(str(a.get('value', a.get('description', ''))) for a in r['params']['args'])[:300])
        elif m:
            self.events.append(r)

    async def ev(self, expr):
        r = await self.send('Runtime.evaluate', {'expression': expr, 'returnByValue': True, 'awaitPromise': True})
        if 'exceptionDetails' in r:
            raise RuntimeError(f'{expr[:80]}: {r["exceptionDetails"].get("exception", {}).get("description", r["exceptionDetails"])}')
        return r['result'].get('value')

    async def read_after(self, action):
        """Run action, which makes the page queue a read, and wait for it.

        The page's queue is a promise each read replaces, so a new one proves
        the read was queued, and awaiting the queue until it stops changing
        waits for it and anything queued behind it."""
        await self.ev('window.__queue = queue')
        await action()
        for _ in range(200):
            if await self.ev('queue !== window.__queue'):
                break
            await asyncio.sleep(0.02)
        else:
            raise TimeoutError('the page queued no read')
        await asyncio.wait_for(self.ev('(async () => { let q; do { q = queue; await q; } while (q !== queue); })()'), 120)

    async def drop(self, paths, x=200, y=150):
        async def act():
            data = {'items': [], 'files': paths, 'dragOperationsMask': 1}
            for kind in ('dragEnter', 'dragOver', 'drop'):
                await self.send('Input.dispatchDragEvent', {'type': kind, 'x': x, 'y': y, 'data': data})
        await self.read_after(act)

    async def choose(self, selector, paths):
        async def act():
            doc = await self.send('DOM.getDocument')
            node = await self.send('DOM.querySelector', {'nodeId': doc['root']['nodeId'], 'selector': selector})
            await self.send('DOM.setFileInputFiles', {'nodeId': node['nodeId'], 'files': paths})
        await self.read_after(act)


ROWS = 'rows.map(r => [r.folder, r.file, r.sheet, r.row, String(r.id), String(r.reviewer), String(r.status), String(r.comment)])'
STATUS = "({text: $('status-text').textContent, error: $('status').classList.contains('error'), shown: $('status').classList.contains('show'),"\
    " files: [...$('status-files').querySelectorAll('li')].map(li => li.textContent), listed: !$('status-files').hidden,"\
    " stop: !$('stop-btn').hidden, busy: $('status').hasAttribute('aria-busy')})"
OPTIONS = "(id => [...$(id).options].map(o => [o.value, o.textContent]))"


async def main():
    failures = []

    def check(ok, what, detail=''):
        print(('ok    ' if ok else 'FAIL  ') + what + ('' if ok else f'\n      {detail}'))
        if not ok:
            failures.append(what)

    version = json.load(urllib.request.urlopen(f'http://127.0.0.1:{CDP}/json/version'))
    browser = await websockets.connect(version['webSocketDebuggerUrl'], max_size=50_000_000)
    await browser.send(json.dumps({'id': 1, 'method': 'Target.createTarget', 'params': {'url': 'about:blank'}}))
    target = json.loads(await browser.recv())['result']['targetId']
    listing = json.load(urllib.request.urlopen(f'http://127.0.0.1:{CDP}/json/list'))
    ws = await websockets.connect([t for t in listing if t['id'] == target][0]['webSocketDebuggerUrl'], max_size=50_000_000)
    page = Page(ws)
    base = tempfile.mkdtemp(prefix='skb-qc-test-')
    try:
        for domain in ('Runtime', 'Page', 'DOM', 'Network'):
            await page.send(f'{domain}.enable')
        await page.send('Network.setCacheDisabled', {'cacheDisabled': True})
        await page.send('Page.setInterceptFileChooserDialog', {'enabled': True})
        await page.send('Emulation.setDeviceMetricsOverride', {'width': 1280, 'height': 900, 'deviceScaleFactor': 1, 'mobile': False})
        await page.send('Page.navigate', {'url': URL})
        for _ in range(200):
            if await page.ev("typeof XLSX !== 'undefined' && document.readyState === 'complete'"):
                break
            await asyncio.sleep(0.05)

        # The SheetJS the page loads is its own copy, and not one with the two published vulnerabilities.
        ver = await page.ev('XLSX.version')
        src = await page.ev("[...document.scripts].map(e => e.getAttribute('src') || '').find(x => /xlsx/.test(x)) || ''")
        check(f'vendors/js/xlsx-{ver}.full.min.js' in src and tuple(int(x) for x in ver.split('.')) >= (0, 20, 2),
              'SheetJS is the vendored copy, 0.20.2 or later (CVE-2023-30533, CVE-2024-22363)', [ver, src])

        # An .xls (BIFF8) workbook, written by the page's own SheetJS.
        xls = base64.b64decode(await page.ev(
            "(() => { const wb = XLSX.utils.book_new();"
            " XLSX.utils.book_append_sheet(wb, XLSX.utils.aoa_to_sheet([['ID', 'QC reviewer', 'QC comment', 'QC status'],"
            " ['p5', 'revB', 'old format', 'OK-PSAR']]), 'Gammal');"
            " return XLSX.write(wb, {bookType: 'biff8', type: 'base64'}); })()"))
        expected = build_tree(base, xls)
        root = os.path.join(base, ROOT)

        # 1. The folder dropped on the zone.
        check(await page.ev("$('workspace').classList.contains('show')") is False, 'the workspace is hidden before anything is read')
        await page.drop([root])
        got = [tuple(r) for r in await page.ev(ROWS)]
        check(got == expected, 'a dropped folder gives every QC row, with its folder, file, sheet and Excel row',
              f'got {got}')
        st = await page.ev(STATUS)
        check(st['text'] == '6 workbook(s) read from 4 folder(s), 12 data point(s) added. 1 gave no data.' and not st['error'],
              'the status counts workbooks and folders, and is no error', st)
        check(st['files'] == [f'{ROOT}/notes.xlsx: no sheet with ID and the three QC columns'] and st['listed'],
              'the workbook without QC columns is listed with its reason', st['files'])
        check(not st['stop'] and not st['busy'], 'Stop is hidden and the status not busy once reading is done', st)
        check(await page.ev("loadedFiles.some(p => /secret|~\\$|\\._/.test(p))") is False,
              'owner files, "._" copies and hidden folders are passed over')
        check(await page.ev("$('workspace').classList.contains('no-folders')") is False
              and await page.ev("getComputedStyle(document.querySelector('th.folder')).display") != 'none',
              'the Folder column is shown')
        cells = await page.ev("[...document.querySelector('#data-body tr').children].slice(0, 4).map(td => td.textContent)")
        check(cells == [f'{ROOT}/PSAR 2024', 'params.xlsx', 'Params', '2'], 'the table row starts with folder, file, sheet, row', cells)

        folders = await page.ev(OPTIONS + "('folder-filter')")
        check(folders == [[ROOT, ROOT], [f'{ROOT}/PSAR 2024', ' PSAR 2024'], [f'{ROOT}/PSAR 2025', ' PSAR 2025'],
                          [f'{ROOT}/PSAR 2025/äldre', '  äldre']],
              'the Folders filter lists the tree, indented', folders)
        files = await page.ev(OPTIONS + "('file-filter')")
        check([v for v, _ in files] == [f'{ROOT}/PSAR 2024/params.xlsx', f'{ROOT}/PSAR 2025/äldre/legacy.xlsm', f'{ROOT}/PSAR 2025/gammal.xls',
                                        f'{ROOT}/PSAR 2025/offset.xlsx', f'{ROOT}/PSAR 2025/params.xlsx']
              and files[0][1] == f'params.xlsx ({ROOT}/PSAR 2024)',
              'two workbooks of one name in two folders are two files, labelled name first', files)
        p1 = await page.ev("(() => { const e = parameterIndex(rows).get('p1'); return [e.files.size, e.why]; })()")
        check(p1 == [2, ['rows disagree on status', '2 files']], 'p1 is flagged as in two files', p1)

        await page.ev(f"(() => {{ const s = $('folder-filter'); [...s.options].forEach(o => o.selected = o.value === {json.dumps(ROOT + '/PSAR 2025')});"
                      " s.dispatchEvent(new Event('change')); })()")
        picked = await page.ev("filtered().map(r => r.id)")
        check(sorted(picked) == ['p1', 'p3', 'p3', 'p4', 'p5', 'p6', 'p7'], 'a folder chosen takes its subfolders along', picked)
        csv_text = await page.ev(
            "(() => { const oc = URL.createObjectURL, orv = URL.revokeObjectURL, ock = HTMLAnchorElement.prototype.click; let blob = null;"
            " URL.createObjectURL = b => { blob = b; return 'blob:stub'; }; URL.revokeObjectURL = () => {}; HTMLAnchorElement.prototype.click = function () {};"
            " try { $('export-btn').click(); } finally { URL.createObjectURL = oc; URL.revokeObjectURL = orv; HTMLAnchorElement.prototype.click = ock; }"
            " return blob.text(); })()")
        table = list(csv.reader(io.StringIO(csv_text.lstrip('﻿')), delimiter=';'))
        check(table[0] == ['Folder', 'File', 'Sheet', 'Excel row', 'ID', 'QC reviewer', 'QC status', 'QC comment']
              and [f'{ROOT}/PSAR 2025', 'offset.xlsx', 'Lower', '4', 'p6', 'revA', 'OK-PSAR', ''] in table and len(table) == 8,
              'the CSV export starts with the Folder column and follows the filter', table[:3])
        await page.ev("(() => { [...$('folder-filter').options].forEach(o => o.selected = false); $('folder-filter').dispatchEvent(new Event('change')); })()")

        # 2. The same folder dropped again, and 3. chosen with the folder picker.
        await page.drop([root])
        st = await page.ev(STATUS)
        check(await page.ev('rows.length') == 12 and '6 had been read before' in st['text'],
              'a folder dropped again replaces its rows', st['text'])
        await page.choose('#folder-input', [root])
        got = [tuple(r) for r in await page.ev(ROWS)]
        check(sorted(got) == sorted(expected), 'the folder picker gives the same rows, still not doubled', got)

        # 4. A workbook added on its own.
        await page.drop([os.path.join(base, 'Lös', 'params.xlsx')])
        lone = await page.ev("rows.filter(r => r.reviewer === 'revE').map(r => [r.folder, r.path])")
        cell = await page.ev("[...document.querySelectorAll('#data-body tr')].find(tr => tr.textContent.includes('revE')).children[0].textContent")
        check(lone == [['', 'params.xlsx']] and cell == '—', 'a workbook dropped on its own has no folder, shown as a dash', [lone, cell])
        p1 = await page.ev("parameterIndex(rows).get('p1').files.size")
        check(p1 == 3, 'and is a third file named params.xlsx', p1)

        # 5. A folder with a broken workbook, and 6. one with no workbook.
        await page.drop([os.path.join(base, 'Trasig')])
        st = await page.ev(STATUS)
        check(st['error'] and len(st['files']) == 1 and st['files'][0].startswith('Trasig/broken.xlsx: could not be read')
              and await page.ev("rows.some(r => r.id === 'p8')"),
              'a workbook that cannot be read is named in an error status, the others still read', st)
        await page.drop([os.path.join(base, 'Tom')])
        st = await page.ev(STATUS)
        check(st['error'] and st['text'] == 'No Excel workbook (.xlsx, .xlsm, .xlsb or .xls) among what was added, only 1 other file(s).',
              'a folder without a workbook says so', st['text'])

        # 7. Stop, and 8. Clear all, while the third of twelve workbooks is read.
        stop_at = ("(button => { const t = $('status-text'); window.__seen = null; const mo = new MutationObserver(() => {"
                   " if (t.textContent.startsWith('Reading 3 of 12')) { mo.disconnect(); window.__seen = !$('stop-btn').hidden; $(button).click(); } });"
                   " mo.observe(t, {childList: true, characterData: true, subtree: true}); })")
        many = os.path.join(base, 'Många')
        await page.ev(stop_at + "('stop-btn')")
        await page.drop([many])
        st = await page.ev(STATUS)
        ids = await page.ev("rows.filter(r => r.reviewer === 'revM').map(r => r.id)")
        check(await page.ev('window.__seen') is True and st['text'].startswith('Stopped after 3 of 12 workbook(s)') and ids == ['m01', 'm02', 'm03'],
              'Stop ends the read and keeps what was read', [st['text'], ids])
        await page.ev(stop_at + "('reset-btn')")
        await page.drop([many])
        left = await page.ev("[rows.length, loadedFiles.length, $('workspace').className, $('status').className]")
        check(left == [0, 0, '', ''], 'Clear all in the middle of a read keeps nothing', left)

        # 9. A drop beside the zone, on the page's header.
        await page.ev("window.__prevented = null; window.addEventListener('drop', e => { window.__prevented = e.defaultPrevented; }, {once: true})")
        await page.drop([root], x=640, y=12)
        check(await page.ev('window.__prevented') is True and await page.ev('rows.length') == 12,
              'a folder dropped outside the zone is taken, and the browser does not open it')

        # 10. The plain file chooser: files have no folder.
        await page.choose('#file-input', [os.path.join(base, 'Lös', 'params.xlsx')])
        check(await page.ev("rows.filter(r => r.folder === '').length") == 1, 'a file chosen with Choose files has no folder')

        # 11. "Choose a folder" opens the folder chooser and the zone the file chooser, by click and by Enter.
        #     The inputs sit inside the zone, so the click a button passes on to its input reaches the zone too.
        async def chooser_after(target, how):
            page.events.clear()
            if how == 'enter':
                await page.ev(f"$({json.dumps(target)}).focus()")
                for kind in ('keyDown', 'keyUp'):
                    await page.send('Input.dispatchKeyEvent', {'type': kind, 'key': 'Enter', 'code': 'Enter', 'windowsVirtualKeyCode': 13,
                                                               **({'text': '\r'} if kind == 'keyDown' else {})})
            else:
                # The zone is clicked beside its buttons, on its label.
                box = await page.ev(f"(() => {{ const el = {'document.querySelector(\'.dz-label\')' if target == 'dropzone' else f'$({json.dumps(target)})'};"
                                    " el.scrollIntoView({block: 'center'}); const b = el.getBoundingClientRect(); return [b.x + b.width / 2, b.y + b.height / 2]; })()")
                for kind in ('mousePressed', 'mouseReleased'):
                    await page.send('Input.dispatchMouseEvent', {'type': kind, 'x': box[0], 'y': box[1], 'button': 'left', 'clickCount': 1})
            await page.ev('new Promise(r => setTimeout(r, 300))')
            opened = [e['params'] for e in page.events if e.get('method') == 'Page.fileChooserOpened']
            if not opened:
                return None
            doc = await page.send('DOM.getDocument')
            ids = {}
            for sel in ('#file-input', '#folder-input'):
                node = await page.send('DOM.querySelector', {'nodeId': doc['root']['nodeId'], 'selector': sel})
                ids[(await page.send('DOM.describeNode', {'nodeId': node['nodeId']}))['node']['backendNodeId']] = sel
            return [ids.get(o.get('backendNodeId')) for o in opened]
        for how in ('click', 'enter'):
            for target, wanted in (('folder-btn', '#folder-input'), ('choose-btn', '#file-input'), ('dropzone', '#file-input')):
                opened = await chooser_after(target, how)
                check(opened == [wanted], f'{how} on #{target} opens only {wanted}', opened)

        # 12. A dropped item with no entry to it, as a mail attachment has, is taken as a file on its own.
        entry = await page.ev(
            "(() => { const wb = XLSX.utils.book_new(); XLSX.utils.book_append_sheet(wb, XLSX.utils.aoa_to_sheet("
            "[['ID', 'QC reviewer', 'QC comment', 'QC status'], ['p9', 'revF', '', 'OK-PSAR']]), 'Params');"
            " window.__attachment = new File([XLSX.write(wb, {bookType: 'xlsx', type: 'array'})], 'bilaga.xlsx');"
            " const dt = new DataTransfer(); dt.items.add(window.__attachment); return dt.items[0].webkitGetAsEntry(); })()")
        await page.read_after(lambda: page.ev(
            "(() => { const dt = new DataTransfer(); dt.items.add(window.__attachment);"
            " document.body.dispatchEvent(new DragEvent('drop', {dataTransfer: dt, bubbles: true, cancelable: true})); })()"))
        got = await page.ev("rows.filter(r => r.id === 'p9').map(r => [r.folder, r.path])")
        check(entry is None and got == [['', 'bilaga.xlsx']], 'a dropped item with no entry is read as a file on its own', [entry, got])

        # 13. No sideways scroll, at a desktop width and on a phone (where fields under 16px make iOS zoom).
        wide = "[document.documentElement.scrollWidth, document.documentElement.clientWidth]"
        check((lambda w: w[0] <= w[1])(await page.ev(wide)), 'no sideways scroll at 1280 px', await page.ev(wide))
        await page.send('Emulation.setDeviceMetricsOverride', {'width': 375, 'height': 548, 'deviceScaleFactor': 2, 'mobile': True})
        await page.send('Emulation.setTouchEmulationEnabled', {'enabled': True, 'maxTouchPoints': 5})
        await page.send('Page.reload', {'ignoreCache': True})
        for _ in range(200):
            if await page.ev("typeof XLSX !== 'undefined' && typeof queue !== 'undefined' && document.readyState === 'complete'"):
                break
            await asyncio.sleep(0.05)
        await page.drop([root])
        phone = await page.ev(wide + ".concat(['id-search', 'folder-filter'].map(id => parseFloat(getComputedStyle($(id)).fontSize)))")
        check(phone[0] <= phone[1] == 375 and min(phone[2:]) >= 16, 'no sideways scroll on a 375 px phone, and its fields are 16 px', phone)

        # 14. In the dark theme every summary card's number stands out from its card (4.5:1 or more).
        contrast = await page.ev(
            "(() => { document.documentElement.setAttribute('data-theme', 'dark');"
            " const c = document.createElement('canvas').getContext('2d', {willReadFrequently: true});"
            " const rgb = s => { c.clearRect(0, 0, 1, 1); c.fillStyle = '#000'; c.fillStyle = s; c.fillRect(0, 0, 1, 1); return [...c.getImageData(0, 0, 1, 1).data].slice(0, 3); };"
            " const lum = a => a.map(v => { v /= 255; return v <= 0.03928 ? v / 12.92 : ((v + 0.055) / 1.055) ** 2.4; }).reduce((s, v, i) => s + v * [0.2126, 0.7152, 0.0722][i], 0);"
            " return [...document.querySelectorAll('.card')].map(card => { const [x, y] = [lum(rgb(getComputedStyle(card.querySelector('.card-value')).color)),"
            " lum(rgb(getComputedStyle(card).backgroundColor))].sort((p, q) => q - p); return Math.round((x + 0.05) / (y + 0.05) * 10) / 10; }); })()")
        check(len(contrast) == 5 and min(contrast) >= 4.5, 'in the dark theme every summary card reads (contrast 4.5:1 or more)', contrast)

        check(not page.errors, 'no exception or console error', page.errors)
    finally:
        shutil.rmtree(base, ignore_errors=True)
        await ws.close()
        await browser.send(json.dumps({'id': 2, 'method': 'Target.closeTarget', 'params': {'targetId': target}}))
        await browser.close()

    print(f'\n{len(failures)} failed' if failures else '\nall passed')
    return 1 if failures else 0


if __name__ == '__main__':
    sys.exit(asyncio.run(main()))
