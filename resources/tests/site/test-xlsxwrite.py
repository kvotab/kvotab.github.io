#!/usr/bin/env python3
"""xlsxwrite.js, the site's Excel writer: a cell comes out in the format it
was written with, and a sheet is one sheet whichever call made it.

    python3 resources/tests/site/test-xlsxwrite.py

xlsxwrite-parts.mjs runs the writer under Node and prints the parts of the
workbook; they are zipped here and read back with openpyxl, which resolves a
cell's style the way a spreadsheet program does. Needs node and openpyxl, not
the server or the browser. With SML_XSD set to the path of a transitional
sml.xsd (ECMA-376 / ISO-IEC 29500-4), the parts are checked against it too.

The faults this was written for (2026-10-04):

  * every format came out one place late: the writer wrote a style for the
    default format after the one standing for it, so a cell named the style
    before its own -- the bold header plain, the cells after it bold;
  * a number format never reached the style, so every number was General;
  * a sheet made by writeData was unknown to the calls that find a sheet by
    name, so a setColumn after it made a second sheet of the same name;
  * XlsxWriter's names for an alignment (valign 'vcenter') and its pattern
    numbers were written as they are, which OOXML does not allow: openpyxl
    would not open the file, and Excel had to repair it.

Exit status is 0 when every check passes.
"""
import io
import json
import os
import subprocess
import sys
import tempfile
import zipfile
from xml.etree import ElementTree as ET

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

failures = []
checks = 0


def check(label, got, want=True):
    global checks
    checks += 1
    ok = got == want
    print(f'{"ok  " if ok else "FAIL"}  {label}' + ('' if ok else f': {got!r} (expected {want!r})'))
    if not ok:
        failures.append(label)


def workbook(case):
    """The parts the writer made for `case`, and the workbook they make."""
    out = subprocess.run(['node', os.path.join(HERE, 'xlsxwrite-parts.mjs'), case],
                         check=True, capture_output=True, text=True).stdout
    parts = json.loads(out)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, text in parts.items():
            z.writestr(name, text)
    buf.seek(0)
    try:
        wb, said = openpyxl.load_workbook(buf), 'opens'
    except Exception as e:  # noqa: BLE001 -- an invalid styles.xml stops openpyxl at the door
        wb, said = None, str(e).split('\n')[0]
    check(f'{case}: the workbook opens', said, 'opens')
    return parts, wb


def formats():
    parts, wb = workbook('formats')
    schema_check('formats', parts)
    if wb is None:
        return
    ws = wb['Formats']
    a = [ws.cell(row=r, column=1) for r in range(1, 9)]
    check('a bold cell is bold', a[0].font.b, True)
    check('and the cell of the format made after it is not', a[1].font.b, False)
    check('a custom number format reaches its cell', a[1].number_format, '0.000E+00')
    check('and so does a built-in one', a[2].number_format, '0.00')
    check('italic, size and colour together', (a[3].font.i, a[3].font.sz, a[3].font.color.rgb),
          (True, 9, 'FF7A6E62'))
    check('with the alignment', a[3].alignment.horizontal, 'center')
    check('a struck-out cell is struck out', a[4].font.strike, True)
    check('a filled, bordered header', (a[5].fill.patternType, a[5].fill.fgColor.rgb, a[5].font.b,
                                        a[5].font.color.rgb, a[5].border.left.style),
          ('solid', 'FF217346', True, 'FFFFFFFF', 'thin'))
    check('the default format is the plain style', (a[6].font.b, a[6].number_format, a[6].style_id), (False, 'General', 0))
    check('and so is no format', (a[7].font.b, a[7].number_format, a[7].style_id), (False, 'General', 0))
    col = ws.column_dimensions['B']
    check('a column gets its width', col.width, 30)
    check('and its format', col.font.b, True)

    # What the counts in styles.xml say, which openpyxl reads past.
    styles = ET.fromstring(parts['xl/styles.xml'])
    xfs = styles.findall('s:cellXfs/s:xf', NS)
    fonts = [ET.tostring(f) for f in styles.findall('s:fonts/s:font', NS)]
    check('one style a format, the default included, and no more', len(xfs), 7)
    check('no font twice', len(fonts), len(set(fonts)))
    sci = xfs[int(ET.fromstring(parts['xl/worksheets/sheet1.xml'])
                  .find('.//s:c[@r="A2"]', NS).get('s'))]
    check('a number format alone shares the default font', (sci.get('fontId'), sci.get('applyFont')), ('0', None))


def schema_check(case, parts):
    """The styles, the workbook and the sheets against the transitional
    schema, where SML_XSD names its sml.xsd (ECMA-376 / ISO-IEC 29500-4)."""
    xsd = os.environ.get('SML_XSD')
    if not xsd or not os.path.exists(xsd):
        print(f'skip  {case}: the schema check (SML_XSD is not the path of a transitional sml.xsd)')
        return
    names = ['xl/styles.xml', 'xl/workbook.xml'] + sorted(n for n in parts if n.startswith('xl/worksheets/sheet'))
    bad = []
    with tempfile.TemporaryDirectory() as tmp:
        for name in names:
            path = os.path.join(tmp, name.replace('/', '_'))
            with open(path, 'w', encoding='utf-8') as fh:
                fh.write(parts[name])
            r = subprocess.run(['xmllint', '--noout', '--nonet', '--schema', xsd, path], capture_output=True, text=True)
            if r.returncode:
                bad.append(name)
    check(f'{case}: every part the schema covers is valid by it', bad, [])


def names():
    # XlsxWriter's names, which a Format takes and OOXML spells otherwise.
    # Written as they were, they made styles.xml invalid (rb's dataset export
    # asks for valign 'vcenter'): openpyxl would not open the file at all,
    # and Excel had to repair it.
    parts, wb = workbook('names')
    schema_check('names', parts)
    if wb is None:
        return
    ws = wb['Names']
    check('XlsxWriter\'s alignment names in OOXML\'s', (ws['A1'].alignment.horizontal, ws['A1'].alignment.vertical),
          ('centerContinuous', 'center'))
    check('and its pattern numbers', (ws['A2'].fill.patternType, ws['A2'].fill.fgColor.rgb), ('solid', 'FFE2EFDA'))


def sheets():
    parts, wb = workbook('sheets')
    schema_check('sheets', parts)
    if wb is None:
        return
    check('a sheet from writeData stays one sheet', wb.sheetnames, ['Settings', 'Second'])
    ws = wb['Settings']
    check('with its data', [[c.value for c in r] for r in ws.iter_rows(min_row=1, max_row=2)], [['a', 'b'], [1, 2]])
    check('the width setColumn gave it', ws.column_dimensions['A'].width, 22)
    check('the cell write added to it', ws['A4'].value, 'after')
    check('and the frozen row', ws.freeze_panes, 'A2')
    names = [s.get('name') for s in ET.fromstring(parts['xl/workbook.xml']).findall('s:sheets/s:sheet', NS)]
    check('the workbook lists each sheet once', names, ['Settings', 'Second'])


def main():
    formats()
    names()
    sheets()
    print(f'\n{checks - len(failures)}/{checks} passed')
    if failures:
        print('FAILED: ' + '; '.join(failures))
        sys.exit(1)


main()
